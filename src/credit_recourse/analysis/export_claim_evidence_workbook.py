from __future__ import annotations
import argparse
import json
from pathlib import Path
from typing import Any
from .claim_evidence_common import canonical_cell, load_yaml
SHEET_SPECS = (('Claims', 'claim_evidence_registry.yaml', 'claims'), ('Sources', 'source_registry.yaml', 'sources'), ('Decisions', 'research_decisions.yaml', 'decisions'), ('Asset Map', 'paper_display_registry.yaml', 'slots'))

def _columns(records: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    for record in records:
        for key in record:
            if key not in columns:
                columns.append(str(key))
    return columns

def run(project_root: Path) -> dict[str, Any]:
    root = project_root.resolve()
    config_root = root / 'src/credit_recourse/configs'
    output = root / 'data/reproduction/claim_evidence/thesis_claim_evidence_map_generated.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise RuntimeError('Generated workbook requires openpyxl from the reproduction lock') from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts: dict[str, int] = {}
    for sheet_name, filename, key in SHEET_SPECS:
        records = load_yaml(config_root / filename)[key]
        columns = _columns(records)
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for record in records:
            sheet.append([canonical_cell(record.get(column)) for column in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(columns, start=1):
            values = [len(str(sheet.cell(row=row, column=index).value or '')) for row in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max(values, default=8) + 2, 10), 45)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        counts[sheet_name] = len(records)
    workbook.save(output)
    result = {'status': 'PASS', 'schema_version': 'claim_map_workbook_v4_1', 'path': output.relative_to(root).as_posix(), 'sheet_counts': counts}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result

def main(argv: list[str] | None=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--project-root', required=True)
    args = parser.parse_args(argv)
    run(Path(args.project_root))
    return 0
if __name__ == '__main__':
    raise SystemExit(main())
