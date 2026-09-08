from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17324D"
TEAL = "0F766E"
BLUE = "2563EB"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E7F6F3"
LIGHT_GRAY = "F3F4F6"
MID_GRAY = "D1D5DB"
DARK_GRAY = "374151"
LIGHT_AMBER = "FEF3C7"
LIGHT_GREEN = "DCFCE7"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D1D5DB")


def safe_token(value: Any) -> str:
    token = re.sub(r'[<>:"/\\|?*]+', "_", str(value or ""))
    token = re.sub(r"\s+", "_", token).strip("_.")
    return token or "UNNAMED"


def scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, list):
        return "; ".join(str(part) for part in value)
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def reviewer_calculation(item: dict[str, Any]) -> str:
    calculation = str(item.get("calculation", ""))
    if calculation == "POLICY_VALUE_SAME_FIRM":
        return "같은 기업의 행동 후 Oracle 점수에서 그 기업의 무행동 Oracle 점수를 뺀 뒤 기업 단위로 집계"
    if calculation == "SOURCE_TABLE_RESHAPE":
        return "선택 실행 산출물에 명시된 조건을 적용하고 논문 분석단위로 행·열을 정리"
    if calculation == "MIXED_ROLE_REGISTRY":
        return "계산 가능한 결과는 실행 산출물에서 계산하고 정의·설정은 별도로 구분"
    return str(item.get("formula") or "선택 실행의 실제 산출물에서 계산")


def is_computed_empirical(item: dict[str, Any]) -> bool:
    if item.get("class") not in {"EMPIRICAL", "MIXED"}:
        return False
    return any(
        output.get("status") == "SELECTOR_COMPUTED" and output.get("rows")
        for output in item.get("selector_outputs", [])
        if isinstance(output, dict)
    )


def new_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    return workbook


def title_block(sheet, title: str, subtitle: str, width: int = 10) -> None:
    width = max(4, width)
    end = get_column_letter(width)
    sheet.merge_cells(f"A1:{end}1")
    sheet["A1"] = title
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(bold=True, color=WHITE, size=14)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells(f"A2:{end}2")
    sheet["A2"] = subtitle
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].font = Font(italic=True, color=DARK_GRAY)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    sheet.sheet_view.showGridLines = False


def style_header(cells: Iterable, fill: str = TEAL) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def style_body(cells: Iterable) -> None:
    for cell in cells:
        cell.font = Font(color=DARK_GRAY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN)


def set_widths(sheet, widths: Sequence[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def append_rows(sheet, start_row: int, rows: Sequence[Sequence[Any]], header: bool = True) -> int:
    for row_offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            sheet.cell(start_row + row_offset, column, scalar(value))
    if rows:
        if header:
            style_header(sheet[start_row][: len(rows[0])])
            if len(rows) > 1:
                style_body(
                    cell
                    for row in sheet.iter_rows(
                        min_row=start_row + 1,
                        max_row=start_row + len(rows) - 1,
                        min_col=1,
                        max_col=len(rows[0]),
                    )
                    for cell in row
                )
        else:
            style_body(
                cell
                for row in sheet.iter_rows(
                    min_row=start_row,
                    max_row=start_row + len(rows) - 1,
                    min_col=1,
                    max_col=max(len(row) for row in rows),
                )
                for cell in row
            )
    return start_row + len(rows)


def add_readme(workbook: Workbook, payload: dict[str, Any], item: dict[str, Any] | None = None) -> None:
    sheet = workbook.create_sheet("README")
    title = "논문 산출물 안내" if item is None else f"{item['item_id']} · {item['caption']}"
    if item is None:
        subtitle = "논문 DOCX의 전체 표·그림과 선택 실행의 실제 계산 결과를 연결한 안내입니다."
    elif is_computed_empirical(item):
        subtitle = "논문 인쇄값을 계산 입력으로 쓰지 않고, 선택 실행의 실제 산출물에서 값을 계산한 Excel입니다."
    elif item.get("class") == "NON_EMPIRICAL":
        subtitle = "경험적 수치 계산 대상이 아니며, 설계·설정 근거와 생산 코드를 확인하는 Excel입니다."
    else:
        subtitle = "원 생산자 snapshot 또는 명시적 계산 자료가 보존되지 않은 항목입니다. 새 계산으로 가장하지 않고 남은 근거와 한계를 표시합니다."
    title_block(
        sheet,
        title,
        subtitle,
        8,
    )
    selected = payload["selected_run"]
    rows: list[list[Any]] = [
        ["항목", "내용"],
        ["선택 실행", selected["run_id"]],
        ["실행 방식", selected["mode"]],
        ["논문 항목 읽기", "논문 DOCX에서 번호·제목·위치만 동적으로 읽음"],
    ]
    if item is not None:
        computed = [
            output
            for output in item.get("selector_outputs", [])
            if output.get("status") == "SELECTOR_COMPUTED"
        ]
        grains = sorted({" + ".join(map(str, output.get("row_keys", []))) for output in computed if output.get("row_keys")})
        filters = [output.get("filters", []) for output in computed if output.get("filters")]
        rows.extend(
            [
                ["논문 위치", f"{item.get('section', '')} · 인쇄면 {item.get('printed_page', '')}"],
                ["관측·집계 단위", "; ".join(grains) or "생산 산출물의 명시된 행 단위"],
                ["선택 조건", json.dumps(filters, ensure_ascii=False) if filters else "추가 필터 없음"],
                ["계산 방법", reviewer_calculation(item)],
                ["생산 코드", item.get("producer", "")],
                ["실제 입력 산출물", "\n".join(source.get("path", "") for source in item.get("source_tables", []))],
            ]
        )
    append_rows(sheet, 4, rows)
    set_widths(sheet, [28, 110])
    sheet.freeze_panes = "A5"


def add_source_data(workbook: Workbook, item: dict[str, Any]) -> list[dict[str, Any]]:
    sheet = workbook.create_sheet("SOURCE")
    computed = is_computed_empirical(item)
    if computed:
        subtitle = "선택 실행에서 읽은 실제 중간·최종 산출물입니다. 아래 CALCULATION이 이 셀들을 직접 참조합니다."
    elif item.get("class") == "NON_EMPIRICAL":
        subtitle = "설계·설정 근거의 경로·해시와 제한된 미리보기입니다. 경험적 계산 입력이 아닙니다."
    else:
        subtitle = "남아 있는 증거의 경로·해시와 제한된 미리보기입니다. 원 생산자 snapshot의 fresh regeneration으로 주장하지 않습니다."
    title_block(
        sheet,
        f"{item['item_id']} · 실제 입력 산출물",
        subtitle,
        12,
    )
    row = 4
    blocks: list[dict[str, Any]] = []
    for table_index, table in enumerate(item.get("source_tables", [])):
        columns = [str(value) for value in table.get("columns", [])]
        data = table.get("rows", []) if isinstance(table.get("rows"), list) else []
        if not computed:
            data = data[:50]
        width = max(6, len(columns))
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=min(width, 20))
        sheet.cell(row, 1, f"입력 {table_index + 1}: {table.get('path', '')}")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        sheet.cell(row, 1, "파일 크기(bytes)")
        sheet.cell(row, 2, table.get("size_bytes"))
        sheet.cell(row, 3, "생산 코드")
        sheet.cell(row, 4, item.get("producer", ""))
        sheet.cell(row, 5, "SHA256")
        sheet.cell(row, 6, table.get("sha256", ""))
        sheet.cell(row, 7, "읽기 상태")
        sheet.cell(row, 8, table.get("read_status", ""))
        if not computed and len(table.get("rows", [])) > len(data):
            sheet.cell(row, 9, "미리보기 행")
            sheet.cell(row, 10, f"{len(data)}/{len(table.get('rows', []))}; 전체 파일은 위 경로와 SHA256으로 확인")
        row += 1
        header_row = row
        for index, name in enumerate(columns, start=1):
            sheet.cell(row, index, name)
        if columns:
            style_header(sheet[row][: len(columns)], BLUE)
        row += 1
        data_start = row
        for values in data:
            for index, value in enumerate(values, start=1):
                sheet.cell(row, index, scalar(value))
            row += 1
        if data and columns:
            style_body(
                cell
                for cells in sheet.iter_rows(
                    min_row=data_start,
                    max_row=row - 1,
                    min_col=1,
                    max_col=len(columns),
                )
                for cell in cells
            )
        blocks.append(
            {
                "source_table_index": table_index,
                "columns": columns,
                "data_start": data_start,
                "data_end": row - 1,
            }
        )
        row += 2
    set_widths(sheet, [28] + [20] * 19)
    sheet.freeze_panes = "A4"
    return blocks


def source_formula(binding: dict[str, Any] | None, blocks: list[dict[str, Any]]) -> str | None:
    if not binding or binding.get("kind") != "DIRECT_SOURCE_CELL":
        return None
    try:
        table_index = int(binding["source_table_index"])
        row_index = int(binding["source_data_row_index"])
        column_index = int(binding["source_column_index"])
        block = blocks[table_index]
    except (KeyError, TypeError, ValueError, IndexError):
        return None
    return f"='SOURCE'!{get_column_letter(column_index + 1)}{block['data_start'] + row_index}"


def add_analysis(workbook: Workbook, item: dict[str, Any], blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sheet = workbook.create_sheet("CALCULATION")
    title_block(
        sheet,
        f"{item['item_id']} · 분석 과정",
        "필터·분석단위와 실제 계산 결과입니다. 수치 셀은 SOURCE를 직접 참조합니다.",
        12,
    )
    method_rows = [
        ["항목", "내용"],
        ["계산 방법", reviewer_calculation(item)],
        ["생산 코드", item.get("producer", "")],
        ["선택 실행", item.get("run_id", "")],
    ]
    append_rows(sheet, 4, method_rows)
    row = 11
    if item.get("reviewer_claim_rows"):
        sheet.cell(row, 1, "이 표·그림과 연결된 본문 핵심 수치")
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        claim_rows = [["수치 ID", "본문 주장", "논문 표시값", "재계산값", "확인 결과", "계산식", "실제 입력 산출물"]]
        claim_rows.extend(
            [
                claim.get("claim_token_id", ""),
                claim.get("thesis_claim", ""),
                claim.get("thesis_value", ""),
                claim.get("recomputed_value", ""),
                claim.get("result", ""),
                claim.get("calculation", ""),
                claim.get("source_outputs", ""),
            ]
            for claim in item["reviewer_claim_rows"]
        )
        row = append_rows(sheet, row, claim_rows) + 2
    if item.get("same_firm_pairing_summaries"):
        rows = [["입력 산출물", "기업 행 수", "결합 방식", "정책가치 정의"]]
        rows.extend(
            [
                summary.get("source_path", ""),
                summary.get("row_count", ""),
                summary.get("pairing_method", ""),
                "동일기업 행동 점수 − 동일기업 무행동 점수",
            ]
            for summary in item["same_firm_pairing_summaries"]
        )
        sheet.cell(row, 1, "동일기업 정책가치 계산")
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row = append_rows(sheet, row + 1, rows) + 2
    if item.get("policy_value_checks"):
        headers = ["정책", "Oracle", "행동 평균점수", "무행동 평균점수", "재계산 정책가치", "생산 산출물 정책가치", "차이"]
        sheet.cell(row, 1, "정책가치 집계")
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        for column, name in enumerate(headers, start=1):
            sheet.cell(row, column, name)
        style_header(sheet[row][: len(headers)], BLUE)
        row += 1
        for check in item["policy_value_checks"]:
            sheet.cell(row, 1, scalar(check.get("policy")))
            sheet.cell(row, 2, scalar(check.get("oracle")))
            sheet.cell(row, 3, scalar(check.get("action_score")))
            sheet.cell(row, 4, scalar(check.get("same_firm_noop_mean_score")))
            sheet.cell(row, 5, f"=C{row}-D{row}")
            sheet.cell(row, 6, scalar(check.get("reported_delta")))
            sheet.cell(row, 7, f"=E{row}-F{row}")
            row += 1
        row += 2
    selector_blocks: list[dict[str, Any]] = []
    for output in item.get("selector_outputs", []):
        if output.get("status") != "SELECTOR_COMPUTED" or not output.get("rows"):
            continue
        sheet.cell(row, 1, f"계산 데이터: {output.get('selector_id', '')}")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_TEAL)
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        metadata = [
            ["입력 파일", "; ".join(output.get("source_paths", []))],
            ["선택 조건", json.dumps(output.get("filters", []), ensure_ascii=False)],
            ["분석 단위", "; ".join(output.get("row_keys", []))],
            ["집계", output.get("aggregation", "")],
        ]
        append_rows(sheet, row, metadata, header=False)
        row += len(metadata)
        columns = [str(value) for value in output.get("columns", [])]
        header_row = row
        for column, name in enumerate(columns, start=1):
            sheet.cell(row, column, name)
        style_header(sheet[row][: len(columns)], BLUE)
        row += 1
        data_start = row
        bindings = output.get("cell_bindings", [])
        for row_index, values in enumerate(output.get("rows", [])):
            for column_index, value in enumerate(values):
                binding = None
                if row_index < len(bindings) and column_index < len(bindings[row_index]):
                    binding = bindings[row_index][column_index]
                formula = source_formula(binding, blocks)
                sheet.cell(row, column_index + 1, formula if formula else scalar(value))
            row += 1
        selector_blocks.append(
            {
                "selector_id": output.get("selector_id", ""),
                "columns": columns,
                "header_row": header_row,
                "data_start": data_start,
                "data_end": row - 1,
                "output": output,
            }
        )
        row += 2
    if not selector_blocks:
        sheet.cell(row, 1, "비계산·보존근거 등록")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        columns = ["근거 구분", "상태", "생산 코드", "실제 근거 경로", "설명"]
        for column, name in enumerate(columns, start=1):
            sheet.cell(row, column, name)
        style_header(sheet[row][: len(columns)], BLUE)
        header_row = row
        row += 1
        data_start = row
        sources = item.get("source_tables", []) or [{}]
        for source in sources:
            evidence_kind = (
                "설계·설정 근거"
                if item.get("class") == "NON_EMPIRICAL"
                else "보존 증거·생산자 snapshot 공백"
            )
            values = [
                evidence_kind,
                item.get("evidence_status", ""),
                item.get("producer", ""),
                source.get("path", ""),
                item.get("contract_note", "") or item.get("selector_reason", ""),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row, column, scalar(value))
            row += 1
        style_body(
            cell
            for cells in sheet.iter_rows(
                min_row=data_start,
                max_row=row - 1,
                min_col=1,
                max_col=len(columns),
            )
            for cell in cells
        )
        selector_blocks.append(
            {
                "selector_id": "NON_EMPIRICAL_REGISTRY" if item.get("class") == "NON_EMPIRICAL" else "PRESERVED_EVIDENCE_GAP",
                "columns": columns,
                "header_row": header_row,
                "data_start": data_start,
                "data_end": row - 1,
                "output": {},
            }
        )
    set_widths(sheet, [36, 60, 20, 20, 24, 70, 80] + [20] * 13)
    sheet.freeze_panes = "A11"
    return selector_blocks


def add_final_table(workbook: Workbook, item: dict[str, Any], selector_blocks: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("PAPER_TABLE")
    title_block(
        sheet,
        f"{item['item_id']} · 논문 표 작업용 결과",
        "CALCULATION에서 계산된 값만 연결한 최종 작업표입니다. Excel에서 정렬·피벗·표시형식만 조정하면 됩니다.",
        12,
    )
    row = 4
    ordered = sorted(
        selector_blocks,
        key=lambda block: (block["selector_id"] != item.get("paper_selector"), block["selector_id"]),
    )
    for block in ordered:
        sheet.cell(row, 1, f"계산 결과: {block['selector_id']}")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
        row += 1
        for index, _ in enumerate(block["columns"], start=1):
            column = get_column_letter(index)
            sheet.cell(row, index, f"='CALCULATION'!{column}{block['header_row']}")
        style_header(sheet[row][: len(block["columns"])])
        row += 1
        for source_row in range(block["data_start"], block["data_end"] + 1):
            for index, _ in enumerate(block["columns"], start=1):
                column = get_column_letter(index)
                sheet.cell(row, index, f"='CALCULATION'!{column}{source_row}")
            row += 1
        row += 2
    set_widths(sheet, [32] + [20] * 19)
    sheet.freeze_panes = "A5"


def select_chart_output(item: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    spec = item.get("chart_spec") or {}
    outputs = [
        output
        for output in item.get("selector_outputs", [])
        if output.get("status") == "SELECTOR_COMPUTED" and output.get("rows")
    ]
    output = next(
        (value for value in outputs if value.get("selector_id") == spec.get("selector_id")),
        outputs[0],
    )
    return spec, output


def add_chart_data(workbook: Workbook, item: dict[str, Any], blocks: list[dict[str, Any]]) -> dict[str, Any]:
    sheet = workbook.create_sheet("CHART_DATA")
    title_block(
        sheet,
        f"{item['item_id']} · 그림 데이터",
        "실제 실행 산출물과 연결된 데이터입니다. 차트는 이 시트의 값을 참조하므로 Excel에서 직접 편집할 수 있습니다.",
        10,
    )
    spec, output = select_chart_output(item)
    columns = [str(value) for value in output.get("columns", [])]
    x_names = spec.get("x") if isinstance(spec.get("x"), list) else [spec.get("x")]
    x_names = [str(value) for value in x_names if value in columns]
    if not x_names:
        x_names = [columns[0]]
    series_names = [str(value) for value in spec.get("series", []) if value in columns]
    if not series_names:
        series_names = [
            name
            for name in output.get("value_columns", [])
            if name in columns
        ][:4]
    if not series_names:
        raise RuntimeError(f"그림 {item['item_id']}에 수치 계열이 없습니다.")
    sheet.cell(4, 1, " × ".join(x_names))
    for index, name in enumerate(series_names, start=2):
        sheet.cell(4, index, name)
    style_header(sheet[4][: 1 + len(series_names)])
    bindings = output.get("cell_bindings", [])
    max_rows = min(len(output.get("rows", [])), int(spec.get("max_rows", 500)))
    for row_index, values in enumerate(output.get("rows", [])[:max_rows], start=5):
        output_row_index = row_index - 5
        category_links: list[str] = []
        for name in x_names:
            source_index = columns.index(name)
            binding = None
            if output_row_index < len(bindings) and source_index < len(bindings[output_row_index]):
                binding = bindings[output_row_index][source_index]
            formula = source_formula(binding, blocks)
            if not formula:
                category_links = []
                break
            category_links.append(formula.removeprefix("="))
        if category_links:
            category = "=" + '&" | "&'.join(category_links)
        else:
            category = " | ".join(str(values[columns.index(name)]) for name in x_names)
        sheet.cell(row_index, 1, category)
        for series_offset, name in enumerate(series_names, start=2):
            source_index = columns.index(name)
            binding = None
            if output_row_index < len(bindings) and source_index < len(bindings[output_row_index]):
                binding = bindings[output_row_index][source_index]
            formula = source_formula(binding, blocks)
            sheet.cell(row_index, series_offset, formula if formula else scalar(values[source_index]))
    set_widths(sheet, [42] + [20] * len(series_names))
    sheet.freeze_panes = "A5"
    return {
        "row_count": max_rows,
        "series_names": series_names,
        "chart_type": str(spec.get("chart_type", "bar")),
    }


def add_chart(workbook: Workbook, item: dict[str, Any], chart_data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("CHART")
    title_block(
        sheet,
        f"{item['item_id']} · {item['caption']}",
        "CHART_DATA를 참조하는 편집 가능한 Excel 차트입니다.",
        14,
    )
    if chart_data["chart_type"] == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
    chart.title = item.get("caption", item.get("title", ""))
    chart.style = 10
    chart.height = 13
    chart.width = 25
    data_end = 4 + chart_data["row_count"]
    data = Reference(
        workbook["CHART_DATA"],
        min_col=2,
        max_col=1 + len(chart_data["series_names"]),
        min_row=4,
        max_row=data_end,
    )
    categories = Reference(workbook["CHART_DATA"], min_col=1, min_row=5, max_row=data_end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend.position = "r"
    sheet.add_chart(chart, "A4")
    set_widths(sheet, [14] * 14)


def add_chart_spec_data(workbook: Workbook, item: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("CHART_DATA")
    title_block(
        sheet,
        f"{item['item_id']} · 그림 근거와 편집 spec",
        "계산 가능한 수치 계열이 없는 그림입니다. 설계·설정 근거 또는 보존 증거의 범위를 그대로 기록합니다.",
        8,
    )
    rows = [
        ["항목", "내용"],
        ["논문 그림", item.get("caption", "")],
        ["구분", item.get("class", "")],
        ["근거 상태", item.get("evidence_status", "")],
        ["생산 코드", item.get("producer", "")],
        ["실제 근거 경로", "\n".join(source.get("path", "") for source in item.get("source_tables", []))],
        ["편집 spec", json.dumps(item.get("chart_spec") or {}, ensure_ascii=False, sort_keys=True)],
        ["한계·설명", item.get("contract_note", "") or item.get("selector_reason", "")],
    ]
    append_rows(sheet, 4, rows)
    set_widths(sheet, [30, 110])


def add_chart_spec(workbook: Workbook, item: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("CHART_SPEC")
    title_block(
        sheet,
        f"{item['item_id']} · 편집 가능한 그림 명세",
        "이 시트의 텍스트와 근거 경로를 사용해 도식 또는 그림을 편집합니다. 경험적 결과가 없는 경우 수치를 임의로 만들지 않습니다.",
        8,
    )
    rows = [
        ["편집 항목", "값"],
        ["제목", item.get("caption", "")],
        ["그림 역할", "설계·설명 도식" if item.get("class") == "NON_EMPIRICAL" else "보존 증거 범위 내 그림"],
        ["구성 근거", item.get("contract_note", "")],
        ["생산 코드", item.get("producer", "")],
        ["입력 경로", "\n".join(source.get("path", "") for source in item.get("source_tables", []))],
        ["재생성 주장", "하지 않음"],
    ]
    append_rows(sheet, 4, rows)
    set_widths(sheet, [30, 110])


def save_workbook(workbook: Workbook, path: Path, expected_sheets: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if workbook.sheetnames != list(expected_sheets):
        raise RuntimeError(f"시트 구성이 다릅니다: {path.name}: {workbook.sheetnames}")
    workbook.save(path)
    reopened = load_workbook(path, read_only=True, data_only=False)
    try:
        if reopened.sheetnames != list(expected_sheets):
            raise RuntimeError(f"저장 후 시트 구성이 다릅니다: {path.name}")
    finally:
        reopened.close()


def build_index(payload: dict[str, Any], output_root: Path) -> None:
    workbook = new_workbook()
    sheet = workbook.create_sheet("OUTPUT_INDEX")
    title_block(
        sheet,
        "논문 표·그림 산출물 목록",
        "논문 DOCX에서 전체 항목을 직접 읽었습니다. 계산 항목은 선택 실행 산출물로, 비계산 항목은 설계·설정 근거로, 보존 공백은 그 한계로 연결했습니다.",
        10,
    )
    headers = ["논문 항목", "종류", "논문 번호", "논문 쪽", "제목", "Excel 산출물", "확인 범위", "계산 방법", "실제 산출물", "생산 코드"]
    rows: list[list[Any]] = [headers]
    for item in payload["items"]:
        computed = is_computed_empirical(item)
        token = "J_UNNUMBERED" if item["item_id"] == "TJ-U" else item.get("thesis_number", "")
        if item.get("kind") == "figure":
            output = f"figures/FIGURE_{safe_token(item.get('thesis_number'))}.xlsx"
        else:
            output = f"tables/TABLE_{safe_token(token)}.xlsx"
        if computed:
            coverage = "선택 실행의 실제 산출물에서 계산"
        elif item.get("class") in {"EMPIRICAL", "MIXED"}:
            coverage = "직접 계산할 원 실행 산출물이 보존되지 않음"
        else:
            coverage = "경험적 계산 대상이 아닌 설계·설명 항목"
        rows.append(
            [
                item.get("item_id"),
                "그림" if item.get("kind") == "figure" else "표",
                item.get("thesis_number"),
                item.get("printed_page"),
                item.get("caption"),
                output,
                coverage,
                reviewer_calculation(item),
                "\n".join(source.get("path", "") for source in item.get("source_tables", [])),
                item.get("producer", ""),
            ]
        )
    append_rows(sheet, 4, rows)
    set_widths(sheet, [14, 9, 12, 10, 56, 44, 34, 76, 88, 72])
    sheet.freeze_panes = "A5"
    add_readme(workbook, payload)
    save_workbook(workbook, output_root / "THESIS_OUTPUT_INDEX.xlsx", ["OUTPUT_INDEX", "README"])


def add_reviewer_calculation_sheets(workbook: Workbook, payload: dict[str, Any]) -> list[str]:
    names: list[str] = []
    for table in payload.get("reviewer_calculation_tables", []):
        name = safe_token(table.get("sheet_name", "ROW_CALCULATIONS"))[:31]
        if not name or name in workbook.sheetnames:
            continue
        columns = [str(value) for value in table.get("columns", [])]
        rows = table.get("rows", []) if isinstance(table.get("rows"), list) else []
        if not columns or not rows:
            continue
        sheet = workbook.create_sheet(name)
        title_block(
            sheet,
            str(table.get("description", name)),
            "선택 실행의 기업별 원행에서 계산했습니다. 아래 계산열은 같은 행의 입력 셀을 직접 참조합니다.",
            max(8, len(columns)),
        )
        append_rows(
            sheet,
            4,
            [
                ["실제 입력·중간 산출물", table.get("path", "")],
                ["행 수", len(rows)],
            ],
            header=False,
        )
        header_row = 7
        for column_index, column_name in enumerate(columns, start=1):
            sheet.cell(header_row, column_index, column_name)
        style_header(sheet[header_row][: len(columns)], BLUE)
        data_start = header_row + 1
        for row_offset, values in enumerate(rows):
            target_row = data_start + row_offset
            for column_index, value in enumerate(values, start=1):
                sheet.cell(target_row, column_index, scalar(value))

        column_map = {column: index + 1 for index, column in enumerate(columns)}
        for target_row in range(data_start, data_start + len(rows)):
            if name == "CANDIDATE_IQL_ROWS":
                for oracle in ("alpha", "beta", "gamma"):
                    result = column_map[f"difference_{oracle}"]
                    left = get_column_letter(column_map[f"candidate_iql_{oracle}"])
                    right = get_column_letter(column_map[f"weakest_component_{oracle}"])
                    sheet.cell(target_row, result, f"={left}{target_row}-{right}{target_row}")
            elif name == "ACTION_ROWS":
                proposed_columns = [
                    column_map[column]
                    for column in columns
                    if column.startswith("proposed__")
                ]
                applied_columns = [
                    column_map[column]
                    for column in columns
                    if column.startswith("applied__")
                ]
                if len(proposed_columns) == 10 and len(applied_columns) == 10:
                    proposed_range = (
                        f"{get_column_letter(min(proposed_columns))}{target_row}:"
                        f"{get_column_letter(max(proposed_columns))}{target_row}"
                    )
                    applied_range = (
                        f"{get_column_letter(min(applied_columns))}{target_row}:"
                        f"{get_column_letter(max(applied_columns))}{target_row}"
                    )
                    proposed_l1 = get_column_letter(column_map["proposed_l1"])
                    applied_l1 = get_column_letter(column_map["applied_l1"])
                    sheet.cell(
                        target_row,
                        column_map["proposed_l1"],
                        f"=SUMPRODUCT(ABS({proposed_range}))",
                    )
                    sheet.cell(
                        target_row,
                        column_map["applied_l1"],
                        f"=SUMPRODUCT(ABS({applied_range}))",
                    )
                    sheet.cell(
                        target_row,
                        column_map["proposed_l1_over_0p75"],
                        f"=--({proposed_l1}{target_row}>0.750000001)",
                    )
                    sheet.cell(
                        target_row,
                        column_map["any_variable_clipping"],
                        f"=--(SUMPRODUCT(--(ABS({proposed_range}-{applied_range})>1E-12))>0)",
                    )
                    sheet.cell(
                        target_row,
                        column_map["applied_l1_at_least_95pct_of_0p75"],
                        f"=--({applied_l1}{target_row}>=0.7125)",
                    )
            elif name == "REPEATABILITY_ROWS":
                left_candidate = get_column_letter(column_map["candidate_id_left"])
                right_candidate = get_column_letter(column_map["candidate_id_right"])
                sheet.cell(
                    target_row,
                    column_map["nearest_candidate_match"],
                    f"=--({left_candidate}{target_row}={right_candidate}{target_row})",
                )
                for oracle in ("alpha", "beta", "gamma"):
                    left = get_column_letter(column_map[f"policy_value_{oracle}_left"])
                    right = get_column_letter(column_map[f"policy_value_{oracle}_right"])
                    sheet.cell(
                        target_row,
                        column_map[f"policy_value_{oracle}_difference"],
                        f"={left}{target_row}-{right}{target_row}",
                    )
            elif name == "C4_C4R_C6_ROWS":
                c4 = get_column_letter(column_map["policy_value_C4_alpha"])
                c4r = get_column_letter(column_map["policy_value_C4R_alpha"])
                c6 = get_column_letter(column_map["policy_value_C6_alpha"])
                self_review = get_column_letter(column_map["self_review_C4R_minus_C4"])
                reference = get_column_letter(column_map["reference_increment_C6_minus_C4R"])
                total = get_column_letter(column_map["total_C6_minus_C4"])
                sheet.cell(target_row, column_map["self_review_C4R_minus_C4"], f"={c4r}{target_row}-{c4}{target_row}")
                sheet.cell(target_row, column_map["reference_increment_C6_minus_C4R"], f"={c6}{target_row}-{c4r}{target_row}")
                sheet.cell(target_row, column_map["total_C6_minus_C4"], f"={c6}{target_row}-{c4}{target_row}")
                sheet.cell(target_row, column_map["identity_error"], f"={total}{target_row}-{self_review}{target_row}-{reference}{target_row}")
            elif name == "E2_E3_ROWS":
                e2 = get_column_letter(column_map["e2_reference_increment_C6_minus_C4R"])
                e3 = get_column_letter(column_map["e3_reference_increment_C6_minus_C4R"])
                sheet.cell(target_row, column_map["e2_minus_e3"], f"={e2}{target_row}-{e3}{target_row}")

        if rows:
            style_body(
                cell
                for cells in sheet.iter_rows(
                    min_row=data_start,
                    max_row=data_start + len(rows) - 1,
                    min_col=1,
                    max_col=len(columns),
                )
                for cell in cells
            )
        set_widths(sheet, [22] * len(columns))
        sheet.freeze_panes = f"A{data_start}"
        names.append(name)
    return names


def build_numeric_claims(payload: dict[str, Any], output_root: Path) -> None:
    workbook = new_workbook()
    sheet = workbook.create_sheet("NUMERIC_CLAIMS")
    title_block(
        sheet,
        "논문 핵심 수치 확인표",
        "논문 표시값은 비교용이며, 재계산값은 선택 실행의 실제 산출물과 명시된 계산식에서 얻었습니다.",
        12,
    )
    headers = ["구분", "논문 쪽", "논문 주장", "논문 표시값", "재계산값", "차이", "확인 결과", "계산식", "실제 산출물", "계산 코드", "실행 ID", "미계산·차이 이유"]
    rows: list[list[Any]] = [headers]
    for claim in payload.get("reviewer_core_claims", []):
        rows.append(
            [
                claim.get("category"), claim.get("printed_page"), f"{claim.get('title', '')}\n\n{claim.get('thesis_claim', '')}",
                claim.get("thesis_values"), claim.get("recomputed_values"), claim.get("differences"),
                claim.get("result"), claim.get("calculation"), claim.get("source_outputs"),
                claim.get("producer"), claim.get("run_id"), claim.get("unavailable_reason"),
            ]
        )
    append_rows(sheet, 4, rows)
    set_widths(sheet, [22, 10, 88, 32, 32, 32, 24, 78, 88, 72, 28, 62])
    sheet.freeze_panes = "A5"

    sample = workbook.create_sheet("SAMPLE_FLOW")
    title_block(sample, "표본 수의 실제 재계산", "각 단계의 분모를 원자료와 선택 실행 산출물에서 다시 계산했습니다.", 10)
    sample_headers = ["선정 단계", "논문 표시값", "실제 재계산값", "차이", "확인 결과", "분모와 의미", "계산 방법", "실제 입력·산출물", "생산 코드", "주의할 점"]
    sample_rows: list[list[Any]] = [sample_headers]
    for row in payload.get("sample_flow", {}).get("rows", []):
        sample_rows.append(
            [
                row.get("stage"), row.get("thesis_value"), row.get("actual_value"), row.get("difference"),
                row.get("result"), row.get("meaning"), row.get("calculation"), row.get("source_outputs"),
                row.get("producer"), row.get("note"),
            ]
        )
    append_rows(sample, 4, sample_rows)
    set_widths(sample, [38, 18, 18, 16, 28, 76, 82, 90, 72, 82])
    sample.freeze_panes = "A5"

    readme = workbook.create_sheet("README")
    title_block(readme, "이 파일을 읽는 법", "교수·심사위원이 논문 핵심 수치를 빠르게 확인하기 위한 요약입니다.", 6)
    append_rows(
        readme,
        4,
        [
            ["항목", "설명"],
            ["정책가치", "같은 기업에서 해당 행동의 Oracle 점수 − 무행동 Oracle 점수"],
            ["논문 표시값", "대조용이며 재계산값을 찾거나 선택하는 입력으로 사용하지 않음"],
            [
                "기업별 계산행",
                "CANDIDATE_IQL_ROWS, ACTION_ROWS, REPEATABILITY_ROWS, C4_C4R_C6_ROWS, E2_E3_ROWS에서 입력행과 Excel 계산식을 직접 확인",
            ],
            ["보존되지 않은 산출물", "새 계산값으로 가장하지 않고 계산 불가 사유를 그대로 표시"],
            ["선택 실행", f"{payload['selected_run']['run_id']} ({payload['selected_run']['mode']})"],
        ],
    )
    set_widths(readme, [30, 110])
    calculation_sheets = add_reviewer_calculation_sheets(workbook, payload)
    save_workbook(
        workbook,
        output_root / "NUMERIC_CLAIMS.xlsx",
        ["NUMERIC_CLAIMS", "SAMPLE_FLOW", "README", *calculation_sheets],
    )


def build_item_workbooks(payload: dict[str, Any], output_root: Path) -> tuple[int, int]:
    table_count = 0
    figure_count = 0
    for item in payload["items"]:
        workbook = new_workbook()
        add_readme(workbook, payload, item)
        blocks = add_source_data(workbook, item)
        if item.get("kind") == "figure":
            if is_computed_empirical(item):
                chart_data = add_chart_data(workbook, item, blocks)
                add_chart(workbook, item, chart_data)
                final_sheet = "CHART"
            else:
                add_chart_spec_data(workbook, item)
                add_chart_spec(workbook, item)
                final_sheet = "CHART_SPEC"
            output = output_root / "figures" / f"FIGURE_{safe_token(item.get('thesis_number'))}.xlsx"
            save_workbook(workbook, output, ["README", "SOURCE", "CHART_DATA", final_sheet])
            figure_count += 1
        else:
            selector_blocks = add_analysis(workbook, item, blocks)
            add_final_table(workbook, item, selector_blocks)
            token = "J_UNNUMBERED" if item["item_id"] == "TJ-U" else item.get("thesis_number")
            output = output_root / "tables" / f"TABLE_{safe_token(token)}.xlsx"
            save_workbook(workbook, output, ["README", "SOURCE", "CALCULATION", "PAPER_TABLE"])
            table_count += 1
    return table_count, figure_count


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build thesis Excel outputs from a prepared selected-run payload.")
    parser.add_argument("--payload", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    if payload.get("schema_version") != "thesis_output_build_payload_v11":
        raise RuntimeError(f"지원하지 않는 계산 자료 형식입니다: {payload.get('schema_version')}")
    if len(payload.get("items", [])) != 71:
        raise RuntimeError("논문 DOCX 항목 수가 예상과 다릅니다.")
    output_root = args.output_root.resolve()
    (output_root / "tables").mkdir(parents=True, exist_ok=True)
    (output_root / "figures").mkdir(parents=True, exist_ok=True)
    build_index(payload, output_root)
    build_numeric_claims(payload, output_root)
    table_count, figure_count = build_item_workbooks(payload, output_root)
    expected = payload.get("workbook_plan", {})
    if table_count != expected.get("table_workbooks") or figure_count != expected.get("figure_workbooks"):
        raise RuntimeError(
            f"Excel 수가 계산 계획과 다릅니다: 표 {table_count}/{expected.get('table_workbooks')}, "
            f"그림 {figure_count}/{expected.get('figure_workbooks')}"
        )
    print(f"논문 수치·표·그림용 Excel {2 + table_count + figure_count}개 생성 완료: {output_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
